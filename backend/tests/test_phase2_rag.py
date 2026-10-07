"""Regression tests for Phase 2 document-grounded RAG."""
from io import BytesIO

import pytest

from app.agents.research import ResearchAgent
from app.models.task import TaskCreate
from app.rag.ingestion import extract_document_segments
from app.tools.vector_search import VectorSearchInput


@pytest.mark.asyncio
async def test_pdf_ingestion_preserves_page_numbers():
    from reportlab.pdfgen import canvas

    buf = BytesIO()
    pdf = canvas.Canvas(buf)
    pdf.drawString(72, 720, "Page one evidence")
    pdf.showPage()
    pdf.drawString(72, 720, "Page two evidence")
    pdf.save()

    segments = await extract_document_segments(buf.getvalue(), "application/pdf", "evidence.pdf")
    assert [s["page_number"] for s in segments] == [1, 2]
    assert "Page one evidence" in segments[0]["text"]
    assert "Page two evidence" in segments[1]["text"]


@pytest.mark.asyncio
async def test_docx_ingestion_extracts_document_text():
    from docx import Document

    buf = BytesIO()
    doc = Document()
    doc.add_paragraph("Quarterly revenue increased.")
    doc.add_paragraph("Operating margin improved.")
    doc.save(buf)

    segments = await extract_document_segments(
        buf.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "results.docx",
    )
    assert len(segments) == 1
    assert "Quarterly revenue increased." in segments[0]["text"]
    assert segments[0]["segment_index"] == 1
    assert "page_number" not in segments[0]


def test_vector_search_schema_never_asks_llm_for_user_id():
    schema = VectorSearchInput.model_json_schema()
    assert "user_id" not in schema.get("properties", {})
    parsed = VectorSearchInput(query="battery safety", file_ids=["file-1"])
    assert parsed.file_ids == ["file-1"]


def test_task_create_accepts_explicit_knowledge_scope():
    task = TaskCreate(
        title="Compare reports",
        description="Use only selected private sources.",
        use_knowledge_base=True,
        selected_file_ids=["file-1", "file-2"],
    )
    assert task.use_knowledge_base is True
    assert task.selected_file_ids == ["file-1", "file-2"]


class _Cursor:
    def __init__(self, rows):
        self.rows = rows

    async def to_list(self, length=None):
        return self.rows[:length] if length else self.rows


class _Files:
    def find(self, query, projection=None):
        ids = query.get("_id", {}).get("$in") if isinstance(query.get("_id"), dict) else None
        rows = [{"_id": "file-1", "original_name": "source.pdf"}]
        if ids is not None:
            rows = [row for row in rows if row["_id"] in ids]
        return _Cursor(rows)


class _Chunks:
    async def find_one(self, query, projection=None):
        if query.get("file_id") == "file-1" and query.get("_id") in (None, "chunk-1"):
            return {"_id": "chunk-1"}
        return None


class _Db:
    files = _Files()
    chunks = _Chunks()


@pytest.mark.asyncio
async def test_research_agent_guarantees_selected_rag_context(monkeypatch):
    import app.agents.research as research_module

    captured = {}

    async def fake_retrieve_chunks(**kwargs):
        captured.update(kwargs)
        return [{
            "chunk_id": "chunk-1",
            "file_id": "file-1",
            "text": "Private evidence from the uploaded PDF.",
            "score": 0.91,
            "metadata": {"original_name": "source.pdf", "page_number": 4},
        }]

    async def fake_agentic_loop(*, messages, available_tools, max_iterations):
        captured["prompt"] = messages[0]["content"]
        captured["tools"] = available_tools
        return {
            "content": '{"summary":"Grounded summary","key_findings":["Private evidence"],"citations":[],"confidence":0.9,"gaps":[]}',
            "tokens": {"prompt": 10, "completion": 10, "total": 20},
            "cost_usd": 0.001,
        }

    monkeypatch.setattr(research_module, "get_db", lambda: _Db())
    monkeypatch.setattr(research_module, "get_redis", lambda: None)
    monkeypatch.setattr(research_module, "retrieve_chunks", fake_retrieve_chunks)

    agent = ResearchAgent(task_id="task-1", step_id="step-1", user_id="user-1")
    monkeypatch.setattr(agent, "run_agentic_loop", fake_agentic_loop)

    output = await agent.run({
        "step_description": "Summarize the uploaded evidence",
        "use_knowledge_base": True,
        "selected_file_ids": ["file-1"],
    })

    assert captured["file_ids"] == ["file-1"]
    assert "Private evidence from the uploaded PDF." in captured["prompt"]
    assert "vector_search" not in captured["tools"]
    assert output["rag_sources"][0]["page_number"] == 4
    local = [c for c in output["citations"] if c.get("type") == "file"]
    assert local and local[0]["file_id"] == "file-1"
    assert local[0]["page_number"] == 4
    assert local[0]["verified"] is True
